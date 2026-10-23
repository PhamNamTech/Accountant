import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

bom_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units)_backup.xlsx'
wb = openpyxl.load_workbook(bom_file, data_only=True)

# Let's inspect DHRM38A6001 BOM items
ws = wb['DHRM38A6001']
for r in range(143, 215):
    name = ws.cell(r, 3).value
    code = ws.cell(r, 4).value
    dvt = ws.cell(r, 6).value
    qty = ws.cell(r, 11).value
    xk = ws.cell(r, 12).value
    note = ws.cell(r, 13).value
    if name:
        print(f'DHRM R{r:3d} | Code: {str(code):15s} | DVT: {str(dvt):5s} | ĐM: {str(qty):>6} | XK: {str(xk):>6} | Name: {str(name)[:35]} | Note: {str(note) if note else ""}')
