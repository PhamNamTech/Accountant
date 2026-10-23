import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

bom_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units).xlsx'
wb = openpyxl.load_workbook(bom_file, data_only=True)

# Check sheets 'Đơn vị chính' and 'đơn vị phụ' in BOM file
ws_main = wb['Đơn vị chính']
ws_sub = wb['đơn vị phụ']
print(f'BOM file - Đơn vị chính: {ws_main.max_row} rows')
print(f'BOM file - đơn vị phụ: {ws_sub.max_row} rows')

# Check THCP codes in BOM file's 'Đơn vị chính'
thcp_codes = set()
for r in range(6, ws_main.max_row):
    thcp = ws_main.cell(r, 21).value
    if thcp:
        thcp_codes.add(thcp)

print('THCP codes in BOM Đơn vị chính:', sorted(list(thcp_codes)))
