import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

wb = openpyxl.load_workbook(r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units)_backup.xlsx', data_only=True)
ws_main = wb['Đơn vị chính']
ws_sub = wb['đơn vị phụ']

print('=== ALL NHẬP KHO NVL THỪA (Đơn vị chính) ===')
for r in range(6, ws_main.max_row):
    sl_nhap = ws_main.cell(r, 13).value
    if sl_nhap and sl_nhap > 0:
        c_code = ws_main.cell(r, 3).value
        name = ws_main.cell(r, 4).value
        so_ct = ws_main.cell(r, 8).value
        dien_giai = str(ws_main.cell(r, 9).value or '')
        dvt = ws_main.cell(r, 10).value
        tk_du = str(ws_main.cell(r, 20).value or '')
        thcp = str(ws_main.cell(r, 21).value or '')
        kmcp = str(ws_main.cell(r, 23).value or '')
        
        # Check reasons for nhap
        if 'thừa' in dien_giai.lower() or 'dư' in dien_giai.lower() or tk_du == '154':
            print(f'R{r}: Code={c_code}, Name={name[:30]}, SL={sl_nhap} {dvt}, THCP={thcp}, TK={tk_du}, Diễn giải="{dien_giai}"')

print('\n=== ALL NHẬP KHO NVL THỪA (đơn vị phụ) ===')
for r in range(6, ws_sub.max_row):
    sl_nhap = ws_sub.cell(r, 13).value
    if sl_nhap and sl_nhap > 0:
        c_code = ws_sub.cell(r, 3).value
        name = ws_sub.cell(r, 4).value
        dien_giai = str(ws_sub.cell(r, 9).value or '')
        dvt = ws_sub.cell(r, 10).value
        tk_du = str(ws_sub.cell(r, 20).value or '')
        thcp = str(ws_sub.cell(r, 21).value or '')
        if 'thừa' in dien_giai.lower() or 'dư' in dien_giai.lower() or tk_du == '154':
            print(f'Sub R{r}: Code={c_code}, Name={name[:30]}, SL={sl_nhap} {dvt}, THCP={thcp}, TK={tk_du}, Diễn giải="{dien_giai}"')
