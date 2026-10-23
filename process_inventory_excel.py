import sys
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

excel_file = r'F:\Accountant\data\So_chi_tiet_vat_tu_hang_hoa.xlsx'
wb = openpyxl.load_workbook(excel_file)

ws1 = wb['đơn vị chính ']
ws2 = wb['đơn vị chuyển đổi']

# 1. Identify transactions and last row for each code in ws1
codes_data_ws1 = {}
for r in range(6, ws1.max_row):
    code = ws1.cell(r, 3).value
    dg = str(ws1.cell(r, 9).value or '').strip()
    nhap_sl = ws1.cell(r, 13).value or 0
    nhap_gt = ws1.cell(r, 14).value or 0
    xuat_sl = ws1.cell(r, 15).value or 0
    xuat_gt = ws1.cell(r, 16).value or 0
    
    if code and code != 'Tổng cộng':
        if code not in codes_data_ws1:
            codes_data_ws1[code] = {
                'first_r': r,
                'last_r': r,
                'has_trans': False,
                'tot_nhap_sl': 0,
                'tot_nhap_gt': 0,
                'tot_xuat_sl': 0,
                'tot_xuat_gt': 0,
            }
        codes_data_ws1[code]['last_r'] = r
        if dg != 'Số dư đầu kỳ':
            codes_data_ws1[code]['tot_nhap_sl'] += nhap_sl
            codes_data_ws1[code]['tot_nhap_gt'] += nhap_gt
            codes_data_ws1[code]['tot_xuat_sl'] += xuat_sl
            codes_data_ws1[code]['tot_xuat_gt'] += xuat_gt
            if (nhap_sl > 0 or nhap_gt > 0 or xuat_sl > 0 or xuat_gt > 0):
                codes_data_ws1[code]['has_trans'] = True

# 2. Identify transactions and last row for each code in ws2
codes_data_ws2 = {}
for r in range(6, ws2.max_row):
    code = ws2.cell(r, 3).value
    dg = str(ws2.cell(r, 9).value or '').strip()
    nhap_sl = ws2.cell(r, 13).value or 0
    nhap_gt = ws2.cell(r, 14).value or 0
    xuat_sl = ws2.cell(r, 15).value or 0
    xuat_gt = ws2.cell(r, 16).value or 0
    
    if code and code != 'Tổng cộng':
        if code not in codes_data_ws2:
            codes_data_ws2[code] = {
                'first_r': r,
                'last_r': r,
                'tot_nhap_sl': 0,
                'tot_nhap_gt': 0,
                'tot_xuat_sl': 0,
                'tot_xuat_gt': 0,
            }
        codes_data_ws2[code]['last_r'] = r
        if dg != 'Số dư đầu kỳ':
            codes_data_ws2[code]['tot_nhap_sl'] += nhap_sl
            codes_data_ws2[code]['tot_nhap_gt'] += nhap_gt
            codes_data_ws2[code]['tot_xuat_sl'] += xuat_sl
            codes_data_ws2[code]['tot_xuat_gt'] += xuat_gt

# 3. Filter ONLY active codes (has_trans == True)
# "giờ tôi muốn chỉ lấy số tồn kho của mấy món có nhập xuất thôi. Ko lấy mấy món không có nhập xuất"
active_codes = [c for c, d in codes_data_ws1.items() if d['has_trans']]

# 4. Sort codes by numeric tail (001, 002, 003, ...)
def sort_key(code):
    m = re.search(r'(\d+)$', str(code))
    if m:
        return (0, int(m.group(1)), str(code))
    else:
        return (1, 0, str(code))

sorted_codes = sorted(active_codes, key=sort_key)
print(f'Total active codes to process: {len(sorted_codes)} (Excluded {len(codes_data_ws1) - len(sorted_codes)} inactive codes)')

# Define common styles
font_family = 'Times New Roman'
header_font = Font(name=font_family, size=10, bold=True)
title_font = Font(name=font_family, size=14, bold=True)
sub_font = Font(name=font_family, size=10, italic=True)
regular_font = Font(name=font_family, size=10)
bold_font = Font(name=font_family, size=10, bold=True)
italic_font = Font(name=font_family, size=10, italic=True)

header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid') # Soft blue-gray
sub_row_fill = PatternFill(start_color='F9FAFC', end_color='F9FAFC', fill_type='solid') # Very soft tint for sub units
total_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid') # Soft yellow for totals

thin_border_side = Side(border_style='thin', color='B0B0B0')
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
double_bottom_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style='double', color='000000'))

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')

num_fmt_qty = '#,##0.00_);(#,##0.00);"-";@'
num_fmt_val = '#,##0_);(#,##0);"-";@'
num_fmt_price = '#,##0.00_);(#,##0.00);"-";@'

# ==========================================
# SHEET 1: Tong_hop_ton_kho (Gọn gàng, lọc chỉ các món có nhập xuất)
# ==========================================
sheet_name_summary = 'Tong_hop_ton_kho'
if sheet_name_summary in wb.sheetnames:
    del wb[sheet_name_summary]
ws_sum = wb.create_sheet(title=sheet_name_summary, index=0)

# Titles
ws_sum['A1'] = 'BẢNG TỔNG HỢP TỒN KHO VẬT TƯ HÀNG HÓA (ĐƠN VỊ CHÍNH & ĐƠN VỊ PHỤ)'
ws_sum['A1'].font = title_font
ws_sum['A2'] = 'Kho: Kho Nguyên vật liệu | Phạm vi: Chỉ lấy các mã có phát sinh nhập/xuất trong kỳ | Sắp xếp: Theo số đuôi mã hàng (001, 002, 003...) | Đến ngày 31/10/2026'
ws_sum['A2'].font = sub_font

headers_sum = [
    ('STT', 6, align_center),
    ('Mã kho', 10, align_center),
    ('Tên kho', 22, align_left),
    ('Mã hàng', 15, align_left),
    ('Tên hàng', 38, align_left),
    ('Phân loại ĐVT', 16, align_center),
    ('ĐVT', 10, align_center),
    ('ĐVT chính', 12, align_center),
    ('Đơn giá tồn', 14, align_right),
    ('Số lượng tồn', 14, align_right),
    ('Giá trị tồn (VND)', 18, align_right),
    ('TK Kho', 10, align_center),
    ('Ghi chú phát sinh trong kỳ', 36, align_left),
]

for col_idx, (h_title, h_width, h_align) in enumerate(headers_sum, 1):
    c = ws_sum.cell(4, col_idx, h_title)
    c.font = header_font
    c.fill = header_fill
    c.alignment = align_center
    c.border = thin_border
    col_letter = get_column_letter(col_idx)
    ws_sum.column_dimensions[col_letter].width = h_width

current_row = 5
stt = 1

for code in sorted_codes:
    d1 = codes_data_ws1[code]
    r1 = d1['last_r']
    ma_kho = ws1.cell(r1, 1).value
    ten_kho = ws1.cell(r1, 2).value
    ten_hang = ws1.cell(r1, 4).value
    tk_kho = ws1.cell(r1, 19).value or '152'
    
    # 1. Main unit row
    dvt1 = ws1.cell(r1, 10).value
    dvt_chinh1 = ws1.cell(r1, 11).value
    don_gia1 = ws1.cell(r1, 12).value
    sl_ton1 = ws1.cell(r1, 17).value or 0
    gt_ton1 = ws1.cell(r1, 18).value or 0
    
    note_main = f'Phát sinh: Nhập {d1["tot_nhap_sl"]:g} {dvt1}, Xuất {d1["tot_xuat_sl"]:g} {dvt1}'
    if sl_ton1 == 0:
        note_main += ' (Đã xuất hết)'
        
    row_cells1 = [
        (stt, align_center, None, regular_font),
        (ma_kho, align_center, None, regular_font),
        (ten_kho, align_left, None, regular_font),
        (code, align_left, None, bold_font),
        (ten_hang, align_left, None, regular_font),
        ('Đơn vị chính', align_center, None, bold_font),
        (dvt1, align_center, None, regular_font),
        (dvt_chinh1, align_center, None, regular_font),
        (don_gia1, align_right, num_fmt_price, regular_font),
        (sl_ton1, align_right, num_fmt_qty, bold_font),
        (gt_ton1, align_right, num_fmt_val, bold_font),
        (tk_kho, align_center, None, regular_font),
        (note_main, align_left, None, regular_font),
    ]
    
    for col_idx, (val, al, nfmt, fn) in enumerate(row_cells1, 1):
        c = ws_sum.cell(current_row, col_idx, val)
        c.font = fn
        c.alignment = al
        c.border = thin_border
        if nfmt:
            c.number_format = nfmt
            
    current_row += 1
    
    # 2. Secondary unit row (if present in ws2)
    if code in codes_data_ws2:
        d2 = codes_data_ws2[code]
        r2 = d2['last_r']
        dvt2 = ws2.cell(r2, 10).value
        dvt_chinh2 = ws2.cell(r2, 11).value
        don_gia2 = ws2.cell(r2, 12).value
        sl_ton2 = ws2.cell(r2, 17).value or 0
        gt_ton2 = ws2.cell(r2, 18).value or 0
        
        note_sub = f'ĐV chuyển đổi (Nhập: {d2["tot_nhap_sl"]:g} {dvt2}, Xuất: {d2["tot_xuat_sl"]:g} {dvt2})'
        if sl_ton2 == 0:
            note_sub += ' (Đã xuất hết)'
            
        row_cells2 = [
            ('', align_center, None, regular_font),
            (ma_kho, align_center, None, regular_font),
            (ten_kho, align_left, None, regular_font),
            (code, align_left, None, regular_font),
            (ten_hang, align_left, None, regular_font),
            ('Đơn vị phụ', align_center, None, italic_font),
            (dvt2, align_center, None, bold_font),
            (dvt_chinh2, align_center, None, regular_font),
            (don_gia2, align_right, num_fmt_price, regular_font),
            (sl_ton2, align_right, num_fmt_qty, bold_font),
            (gt_ton2, align_right, num_fmt_val, regular_font),
            (tk_kho, align_center, None, regular_font),
            (note_sub, align_left, None, italic_font),
        ]
        
        for col_idx, (val, al, nfmt, fn) in enumerate(row_cells2, 1):
            c = ws_sum.cell(current_row, col_idx, val)
            c.font = fn
            c.alignment = al
            c.border = thin_border
            c.fill = sub_row_fill
            if nfmt:
                c.number_format = nfmt
                
        current_row += 1
        
    stt += 1

# Total row for ws_sum
ws_sum.cell(current_row, 1, 'Tổng cộng').font = bold_font
ws_sum.cell(current_row, 1).alignment = align_center
ws_sum.cell(current_row, 1).fill = total_fill

for c_idx in range(2, len(headers_sum) + 1):
    c = ws_sum.cell(current_row, c_idx)
    c.fill = total_fill
    c.border = double_bottom_border
    c.font = bold_font

# Sum of Qty (Primary unit only)
c_sum_qty = ws_sum.cell(current_row, 10)
c_sum_qty.value = f'=SUMIF(F5:F{current_row-1}, "Đơn vị chính", J5:J{current_row-1})'
c_sum_qty.number_format = num_fmt_qty
c_sum_qty.alignment = align_right

# Sum of Value (Primary unit only, avoiding double-counting)
c_sum_val = ws_sum.cell(current_row, 11)
c_sum_val.value = f'=SUMIF(F5:F{current_row-1}, "Đơn vị chính", K5:K{current_row-1})'
c_sum_val.number_format = num_fmt_val
c_sum_val.alignment = align_right

ws_sum.cell(current_row, 13, 'Tổng GT theo ĐV chính').font = italic_font
ws_sum.auto_filter.ref = f'A4:M{current_row-1}'


# ==========================================
# SHEET 2: Ton_kho_24cot (24 cột chuẩn MISA, lọc chỉ các món có nhập xuất)
# ==========================================
sheet_name_24 = 'Ton_kho_24cot'
if sheet_name_24 in wb.sheetnames:
    del wb[sheet_name_24]
ws_24 = wb.create_sheet(title=sheet_name_24, index=1)

# Header block matching ws1 exactly
ws_24['A1'] = 'SỔ CHI TIẾT VẬT TƯ HÀNG HÓA - TỔNG HỢP TỒN KHO CUỐI KỲ (CÁC MÓN CÓ NHẬP XUẤT)'
ws_24['A1'].font = title_font
ws_24['A2'] = ws1['A2'].value or 'Kho: Kho Nguyên vật liệu, Từ ngày 29/05/2026 đến ngày 31/10/2026'
ws_24['A2'].font = sub_font

for col_idx in range(1, 25):
    col_letter = get_column_letter(col_idx)
    ws_24.column_dimensions[col_letter].width = ws1.column_dimensions[col_letter].width

# Copy row 4 and row 5 headers
for col_idx in range(1, 25):
    c4 = ws_24.cell(4, col_idx, ws1.cell(4, col_idx).value)
    c4.font = header_font
    c4.fill = header_fill
    c4.alignment = align_center
    c4.border = thin_border
    
    c5 = ws_24.cell(5, col_idx, ws1.cell(5, col_idx).value)
    c5.font = header_font
    c5.fill = header_fill
    c5.alignment = align_center
    c5.border = thin_border

# Merged cells in headers
merged_24 = [
    'A4:A5', 'B4:B5', 'C4:C5', 'D4:D5', 'E4:E5', 'F4:F5', 'G4:G5', 'H4:H5', 'I4:I5',
    'J4:J5', 'K4:K5', 'L4:L5', 'M4:N4', 'O4:P4', 'Q4:R4',
    'S4:S5', 'T4:T5', 'U4:U5', 'V4:V5', 'W4:W5', 'X4:X5'
]
for m in merged_24:
    ws_24.merge_cells(m)

curr_r_24 = 6

for code in sorted_codes:
    d1 = codes_data_ws1[code]
    r1 = d1['last_r']
    
    row_vals1 = [
        ws1.cell(r1, 1).value,      # 1: Mã kho
        ws1.cell(r1, 2).value,      # 2: Tên kho
        code,                       # 3: Mã hàng
        ws1.cell(r1, 4).value,      # 4: Tên hàng
        ws1.cell(r1, 5).value,      # 5: Mô tả
        '31/10/2026',               # 6: Ngày hạch toán
        None,                       # 7: Ngày chứng từ
        None,                       # 8: Số chứng từ
        'Đơn vị chính',             # 9: Diễn giải
        ws1.cell(r1, 10).value,     # 10: ĐVT
        ws1.cell(r1, 11).value,     # 11: ĐVT chính
        ws1.cell(r1, 12).value,     # 12: Đơn giá
        d1['tot_nhap_sl'],          # 13: Nhập SL (Tổng phát sinh)
        d1['tot_nhap_gt'],          # 14: Nhập GT (Tổng phát sinh)
        d1['tot_xuat_sl'],          # 15: Xuất SL (Tổng phát sinh)
        d1['tot_xuat_gt'],          # 16: Xuất GT (Tổng phát sinh)
        ws1.cell(r1, 17).value or 0,# 17: Tồn SL
        ws1.cell(r1, 18).value or 0,# 18: Tồn GT
        ws1.cell(r1, 19).value or '152', # 19: TK Kho
        None,                       # 20: TK đối ứng
        None,                       # 21: Mã THCP
        None,                       # 22: Tên THCP
        None,                       # 23: Mã KMCP
        None                        # 24: Tên KMCP
    ]
    
    for c_idx, val in enumerate(row_vals1, 1):
        cell = ws_24.cell(curr_r_24, c_idx, val)
        cell.font = bold_font if c_idx in (3, 9, 17, 18) else regular_font
        cell.border = thin_border
        if c_idx in (1, 6, 7, 8, 10, 11, 19, 20):
            cell.alignment = align_center
        elif c_idx in (12, 13, 14, 15, 16, 17, 18):
            cell.alignment = align_right
            if c_idx in (12, 13, 15, 17):
                cell.number_format = num_fmt_qty
            else:
                cell.number_format = num_fmt_val
        else:
            cell.alignment = align_left
            
    curr_r_24 += 1
    
    if code in codes_data_ws2:
        d2 = codes_data_ws2[code]
        r2 = d2['last_r']
        row_vals2 = [
            ws2.cell(r2, 1).value,      # 1: Mã kho
            ws2.cell(r2, 2).value,      # 2: Tên kho
            code,                       # 3: Mã hàng
            ws2.cell(r2, 4).value,      # 4: Tên hàng
            ws2.cell(r2, 5).value,      # 5: Mô tả
            '31/10/2026',               # 6: Ngày hạch toán
            None,                       # 7: Ngày chứng từ
            None,                       # 8: Số chứng từ
            'Đơn vị phụ',               # 9: Diễn giải
            ws2.cell(r2, 10).value,     # 10: ĐVT
            ws2.cell(r2, 11).value,     # 11: ĐVT chính
            ws2.cell(r2, 12).value,     # 12: Đơn giá
            d2['tot_nhap_sl'],          # 13: Nhập SL
            d2['tot_nhap_gt'],          # 14: Nhập GT
            d2['tot_xuat_sl'],          # 15: Xuất SL
            d2['tot_xuat_gt'],          # 16: Xuất GT
            ws2.cell(r2, 17).value or 0,# 17: Tồn SL
            ws2.cell(r2, 18).value or 0,# 18: Tồn GT
            ws2.cell(r2, 19).value or '152', # 19: TK Kho
            None,                       # 20: TK đối ứng
            None,                       # 21: Mã THCP
            None,                       # 22: Tên THCP
            None,                       # 23: Mã KMCP
            None                        # 24: Tên KMCP
        ]
        for c_idx, val in enumerate(row_vals2, 1):
            cell = ws_24.cell(curr_r_24, c_idx, val)
            cell.font = bold_font if c_idx in (10, 17) else (italic_font if c_idx == 9 else regular_font)
            cell.border = thin_border
            cell.fill = sub_row_fill
            if c_idx in (1, 6, 7, 8, 10, 11, 19, 20):
                cell.alignment = align_center
            elif c_idx in (12, 13, 14, 15, 16, 17, 18):
                cell.alignment = align_right
                if c_idx in (12, 13, 15, 17):
                    cell.number_format = num_fmt_qty
                else:
                    cell.number_format = num_fmt_val
            else:
                cell.alignment = align_left
                
        curr_r_24 += 1

# Total row for ws_24
ws_24.cell(curr_r_24, 1, 'Tổng cộng').font = bold_font
ws_24.cell(curr_r_24, 1).alignment = align_center

for c_idx in range(1, 25):
    c = ws_24.cell(curr_r_24, c_idx)
    c.fill = total_fill
    c.border = double_bottom_border
    c.font = bold_font

# Total Nhập, Xuất, Tồn (by Đơn vị chính)
ws_24.cell(curr_r_24, 13, f'=SUMIF(I6:I{curr_r_24-1}, "Đơn vị chính", M6:M{curr_r_24-1})').number_format = num_fmt_qty
ws_24.cell(curr_r_24, 14, f'=SUMIF(I6:I{curr_r_24-1}, "Đơn vị chính", N6:N{curr_r_24-1})').number_format = num_fmt_val
ws_24.cell(curr_r_24, 15, f'=SUMIF(I6:I{curr_r_24-1}, "Đơn vị chính", O6:O{curr_r_24-1})').number_format = num_fmt_qty
ws_24.cell(curr_r_24, 16, f'=SUMIF(I6:I{curr_r_24-1}, "Đơn vị chính", P6:P{curr_r_24-1})').number_format = num_fmt_val
ws_24.cell(curr_r_24, 17, f'=SUMIF(I6:I{curr_r_24-1}, "Đơn vị chính", Q6:Q{curr_r_24-1})').number_format = num_fmt_qty
ws_24.cell(curr_r_24, 18, f'=SUMIF(I6:I{curr_r_24-1}, "Đơn vị chính", R6:R{curr_r_24-1})').number_format = num_fmt_val

ws_24.auto_filter.ref = f'A4:X{curr_r_24-1}'

# Save to original file
wb.save(excel_file)
print(f'Successfully updated: {excel_file}')

# Also save standalone copy
standalone_file = r'F:\Accountant\data\Tong_hop_ton_kho_cuoi_ky.xlsx'
wb.save(standalone_file)
print(f'Successfully saved standalone file: {standalone_file}')
