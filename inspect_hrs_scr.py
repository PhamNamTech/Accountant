import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

bom_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units)_backup.xlsx'
wb = openpyxl.load_workbook(bom_file, data_only=True)

# Inspect HRS20A6003 and SCR18A7501_MC_FSM
for sname in ['HRS20A6003', 'SCR18A7501_MC_FSM']:
    ws = wb[sname]
    start_r = None
    for r in range(1, 200):
        if '1.2' in str(ws.cell(r, 1).value or ''):
            start_r = r
            break
    header_r = start_r + 1
    print(f'=== Sheet {sname} (Header at {header_r}) ===')
    for r in range(header_r + 1, header_r + 100):
        name = ws.cell(r, 3).value
        code = ws.cell(r, 4).value
        dvt = ws.cell(r, 6).value
        qty = ws.cell(r, 11).value
        xk = ws.cell(r, 12).value
        note = ws.cell(r, 13).value
        if name:
            print(f'  R{r:3d} | Code: {str(code):15s} | DVT: {str(dvt):5s} | ĐM: {str(qty):>6} | XK: {str(xk):>6} | Name: {str(name)[:35]} | Note: {str(note) if note else ""}')
    print('\n')
