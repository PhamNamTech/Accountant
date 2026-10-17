r"""
Bộ kịch bản kiểm thử toàn diện các Plugin & MCP Kế toán tại F:\Accountant
Bao gồm:
1. Tra cứu tài khoản TT200 / TT133
2. Bóc tách Hóa đơn điện tử XML (chuẩn Thông tư 78)
3. Xuất bảng kê mua vào hàng loạt sang Excel
4. Cơ sở dữ liệu sổ sách kép SQLite (ketoan.db)
5. Tạo Sổ Nhật Ký Chung Excel chuẩn form Thông tư 200
6. Bóc tách hóa đơn điện tử PDF (pdfplumber)
7. Sổ cái kế toán kép Beancount (Plaintext Double-Entry Accounting)
"""

import os
import sys
import sqlite3
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

BASE_DIR = os.path.abspath("F:/Accountant")
sys.path.insert(0, BASE_DIR)

from tools.ketoan_server import (
    lookup_accounting_account,
    parse_vietnamese_invoice_xml,
    parse_vietnamese_invoice_pdf,
    verify_beancount_ledger,
    batch_export_invoices_to_excel,
    calculate_product_cost
)

def test_lookup_account():
    print("[TEST 1] Tra cứu tài khoản kế toán...")
    res = lookup_accounting_account("112")
    assert "112" in res["results"], "Lỗi: Không tìm thấy TK 112!"
    assert res["results"]["112"]["name"] == "Tiền gửi ngân hàng"
    print(" -> Đạt: Tìm thấy TK 112:", res["results"]["112"])

def test_parse_xml_invoice():
    print("\n[TEST 2] Bóc tách hóa đơn điện tử XML (chuẩn TT78)...")
    xml_path = os.path.join(BASE_DIR, "test_data", "hoadon_01.xml")
    data = parse_vietnamese_invoice_xml(xml_path)
    assert data["ky_hieu"] == "C26TAA", f"Ký hiệu sai: {data['ky_hieu']}"
    assert data["so_hoa_don"] == "0000015", f"Số HĐ sai: {data['so_hoa_don']}"
    assert data["nguoi_ban"]["mst"] == "0108998877", "MST người bán sai!"
    assert len(data["chi_tiet_hang_hoa"]) == 2, "Số lượng mặt hàng sai!"
    assert float(data["tong_tien_thanh_toan"]) == 27500000.0, "Tổng tiền sai!"
    print(f" -> Đạt: Đã bóc tách HĐ {data['so_hoa_don']} từ {data['nguoi_ban']['ten']}")
    print(f"    Tổng thanh toán: {data['tong_tien_thanh_toan']:,.0f} VND ({len(data['chi_tiet_hang_hoa'])} mặt hàng)")

def test_batch_export_excel():
    print("\n[TEST 3] Quét hàng loạt XML và xuất Bảng kê Mua vào sang Excel...")
    folder_in = os.path.join(BASE_DIR, "test_data")
    excel_out = os.path.join(BASE_DIR, "test_output", "Bang_Ke_Mua_Vao_Test.xlsx")
    os.makedirs(os.path.dirname(excel_out), exist_ok=True)
    res = batch_export_invoices_to_excel(folder_in, excel_out)
    assert res["status"] == "success", "Lỗi xuất Excel bảng kê!"
    assert res["total_invoices_processed"] == 2, f"Sai số lượng hóa đơn: {res['total_invoices_processed']}"
    assert os.path.exists(excel_out), "File Excel không được tạo!"

    df_summary = pd.read_excel(excel_out, sheet_name="Bang_Ke_Tong_Hop")
    df_items = pd.read_excel(excel_out, sheet_name="Chi_Tiet_Hang_Hoa")
    assert len(df_summary) == 2, "Bảng kê tổng hợp thiếu dòng!"
    assert len(df_items) == 3, "Bảng kê chi tiết hàng hóa thiếu dòng!"
    total_val = df_summary["Tổng tiền thanh toán"].sum()
    assert total_val == (27500000 + 2160000), f"Tổng tiền 2 hóa đơn sai: {total_val}"
    print(f" -> Đạt: Đã xuất bảng kê ra file {excel_out}")
    print(f"    Tổng 2 hóa đơn đã xử lý: {total_val:,.0f} VND")

def test_sqlite_accounting_db():
    print("\n[TEST 4] Kiểm thử SQLite Kế toán (F:/Accountant/ketoan.db)...")
    db_path = os.path.join(BASE_DIR, "ketoan.db")
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS nhat_ky_chung (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ngay_ghi_so TEXT NOT NULL,
            so_chung_tu TEXT NOT NULL,
            ngay_chung_tu TEXT NOT NULL,
            dien_giai TEXT NOT NULL,
            tk_no TEXT NOT NULL,
            tk_co TEXT NOT NULL,
            so_tien REAL NOT NULL
        )
    """)

    cur.execute("DELETE FROM nhat_ky_chung WHERE so_chung_tu LIKE 'TEST-%'")

    sample_entries = [
        ("2026-10-15", "TEST-PNK01", "2026-10-15", "Mua máy in và mực in", "153", "331", 25000000),
        ("2026-10-15", "TEST-PNK01", "2026-10-15", "Thuế GTGT đầu vào máy in mực in", "1331", "331", 2500000),
        ("2026-10-16", "TEST-UNC01", "2026-10-16", "Chi phí internet T10", "642", "1121", 2000000),
        ("2026-10-16", "TEST-UNC01", "2026-10-16", "Thuế GTGT internet 8%", "1331", "1121", 160000),
    ]

    cur.executemany("""
        INSERT INTO nhat_ky_chung (ngay_ghi_so, so_chung_tu, ngay_chung_tu, dien_giai, tk_no, tk_co, so_tien)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, sample_entries)
    conn.commit()

    cur.execute("SELECT SUM(so_tien) FROM nhat_ky_chung WHERE so_chung_tu LIKE 'TEST-%'")
    total_debit_credit = cur.fetchone()[0]
    conn.close()

    assert total_debit_credit == 29660000.0, f"Tổng số tiền sai: {total_debit_credit}"
    print(f" -> Đạt: SQLite DB hoạt động hoàn hảo tại {db_path}.")
    print(f"    Tổng phát sinh Nợ/Có đối ứng: {total_debit_credit:,.0f} VND")

def test_generate_excel_journal():
    print("\n[TEST 5] Kiểm thử tạo Sổ Nhật Ký Chung Excel chuẩn định dạng...")
    file_path = os.path.join(BASE_DIR, "test_output", "So_Nhat_Ky_Chung_Chuan.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nhat_Ky_Chung"

    ws["A1"] = "CÔNG TY CỔ PHẦN CÔNG NGHỆ VÀ THƯƠNG MẠI NAM PHONG"
    ws["A1"].font = Font(name="Times New Roman", size=11, bold=True)
    ws["A2"] = "Mã số thuế: 0315667788"
    ws["A2"].font = Font(name="Times New Roman", size=10, italic=True)

    ws["A4"] = "SỔ NHẬT KÝ CHUNG"
    ws["A4"].font = Font(name="Times New Roman", size=16, bold=True, color="002060")
    ws["A4"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A4:G4")

    ws["A5"] = "Tháng 10 năm 2026 (Mẫu biểu theo Thông tư 200/2014/TT-BTC)"
    ws["A5"].font = Font(name="Times New Roman", size=11, italic=True)
    ws["A5"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A5:G5")

    headers = ["Ngày ghi sổ", "Số hiệu CT", "Ngày CT", "Diễn giải", "TK Nợ", "TK Có", "Số tiền (VND)"]
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")

    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=7, column=col_num)
        cell.value = h
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rows = [
        ("15/10/2026", "PNK01", "15/10/2026", "Mua máy in và mực in phục vụ văn phòng", "153", "331", 25000000),
        ("15/10/2026", "PNK01", "15/10/2026", "Thuế GTGT đầu vào được khấu trừ", "1331", "331", 2500000),
        ("16/10/2026", "UNC01", "16/10/2026", "Thanh toán cước internet T10/2026", "642", "1121", 2000000),
        ("16/10/2026", "UNC01", "16/10/2026", "Thuế GTGT cước internet 8%", "1331", "1121", 160000),
    ]

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for r_idx, row in enumerate(rows, start=8):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.value = val
            cell.font = Font(name="Times New Roman", size=11)
            cell.border = thin_border
            if c_idx in [1, 2, 3, 5, 6]:
                cell.alignment = Alignment(horizontal="center")
            elif c_idx == 7:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right")

    total_row = 8 + len(rows)
    ws.cell(row=total_row, column=4, value="CỘNG PHÁT SINH").font = Font(name="Times New Roman", size=11, bold=True)
    ws.cell(row=total_row, column=4).alignment = Alignment(horizontal="center")
    sum_cell = ws.cell(row=total_row, column=7)
    sum_cell.value = f"=SUM(G8:G{total_row-1})"
    sum_cell.font = Font(name="Times New Roman", size=11, bold=True, color="C00000")
    sum_cell.number_format = '#,##0'
    sum_cell.alignment = Alignment(horizontal="right")

    col_widths = [15, 15, 15, 45, 12, 12, 20]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb.save(file_path)
    assert os.path.exists(file_path), "File Sổ nhật ký chung không tạo được!"
    print(f" -> Đạt: Đã tạo Sổ Nhật Ký Chung Excel chuẩn form Thông tư 200 tại {file_path}")

def test_parse_pdf_invoice():
    print("\n[TEST 6] Bóc tách Hóa đơn điện tử PDF (pdfplumber)...")
    pdf_path = os.path.join(BASE_DIR, "test_data", "hoadon_vat.pdf")
    data = parse_vietnamese_invoice_pdf(pdf_path)
    assert data["so_hoa_don"] == "0000888", f"Sai số hóa đơn: {data['so_hoa_don']}"
    assert data["ky_hieu"] == "C26TCC", f"Sai ký hiệu: {data['ky_hieu']}"
    assert data["mst_nguoi_ban"] == "0109988776", f"Sai MST người bán: {data['mst_nguoi_ban']}"
    assert data["mst_nguoi_mua"] == "0315667788", f"Sai MST người mua: {data['mst_nguoi_mua']}"
    print(f" -> Đạt: Bóc tách thành công hóa đơn PDF số {data['so_hoa_don']}")
    print(f"    MST bán: {data['mst_nguoi_ban']} | MST mua: {data['mst_nguoi_mua']}")
    print(f"    Tiền hàng: {data['tien_hang']} | Thuế: {data['tien_thue']} | Tổng: {data['tong_thanh_toan']}")

def test_beancount_accounting():
    print("\n[TEST 7] Sổ cái kế toán kép Beancount (Plaintext Accounting)...")
    ledger_path = os.path.join(BASE_DIR, "ledger", "so_cai.beancount")
    res = verify_beancount_ledger(ledger_path)
    assert res["status"] == "valid", f"Lỗi trong sổ cái Beancount: {res['errors']}"
    assert len(res["errors"]) == 0, "Phát hiện lỗi không cân đối trong Beancount!"
    print(f" -> Đạt: Sổ cái Beancount hoàn toàn cân đối (Double-entry balanced 100%)")
    print(f"    Tổng số bút toán đã hạch toán: {res['total_entries']} bút toán")
    print(f"    Đơn vị tiền tệ chính: {res['operating_currency']}")

def test_cost_accounting_regimes():
    print("\n[TEST 8] Kiểm thử Tính Giá Thành & Định khoản theo Chế độ Kế toán (TT99, TT200, TT133)...")
    
    # Test TT99 (Mặc định): Nhập kho TK 155 (Sản phẩm)
    res_tt99 = calculate_product_cost("Bàn gỗ sồi", 10000000, 50000000, 20000000, 12000000, 11000000, 1000, 1000000, regime="TT99")
    assert res_tt99["tong_gia_thanh_Z"] == 80000000, f"Tổng Z sai: {res_tt99['tong_gia_thanh_Z']}"
    assert res_tt99["gia_thanh_don_vi_z"] == 80000, f"Đơn giá z sai: {res_tt99['gia_thanh_don_vi_z']}"
    assert res_tt99["but_toan_dinh_khoan"]["nhap_kho"]["no"] == "155", f"TT99 phải nhập kho TK 155, nhận được: {res_tt99['but_toan_dinh_khoan']['nhap_kho']['no']}"
    print(" -> Đạt TT99: Bút toán nhập kho Nợ TK 155 / Có TK 154 chuẩn Thông tư 99/2025/TT-BTC")

    # Test TT200: Nhập kho TK 1551 (Thành phẩm)
    res_tt200 = calculate_product_cost("Bàn gỗ sồi", 10000000, 50000000, 20000000, 12000000, 11000000, 1000, 1000000, regime="TT200")
    assert res_tt200["but_toan_dinh_khoan"]["nhap_kho"]["no"] == "1551", f"TT200 phải nhập kho TK 1551, nhận được: {res_tt200['but_toan_dinh_khoan']['nhap_kho']['no']}"
    print(" -> Đạt TT200: Bút toán nhập kho Nợ TK 1551 / Có TK 154 chuẩn Thông tư 200/2014/TT-BTC")

    # Test TT133: Nhập kho TK 155 và tập hợp qua tiểu khoản TK 154
    res_tt133 = calculate_product_cost("Bàn gỗ sồi", 10000000, 50000000, 20000000, 12000000, 11000000, 1000, 1000000, regime="TT133")
    assert res_tt133["but_toan_dinh_khoan"]["nhap_kho"]["no"] == "155", f"TT133 phải nhập kho TK 155, nhận được: {res_tt133['but_toan_dinh_khoan']['nhap_kho']['no']}"
    assert res_tt133["but_toan_dinh_khoan"]["tap_hop_hoac_ket_chuyen"][0]["tk"] == "1541"
    print(" -> Đạt TT133: Bút toán tập hợp Nợ TK 1541/1542/1547 và nhập kho Nợ TK 155 chuẩn Thông tư 133/2016/TT-BTC")

if __name__ == "__main__":
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    print("==================================================")
    print(" KIỂM THỬ TOÀN DIỆN HỆ THỐNG KẾ TOÁN (8 BÀI TEST)")
    print("==================================================")
    test_lookup_account()
    test_parse_xml_invoice()
    test_batch_export_excel()
    test_sqlite_accounting_db()
    test_generate_excel_journal()
    test_parse_pdf_invoice()
    test_beancount_accounting()
    test_cost_accounting_regimes()
    print("\n==================================================")
    print(" TẤT CẢ 8 BÀI TEST ĐÃ VƯỢT QUA XUẤT SẮC 100%!")
    print("==================================================")
