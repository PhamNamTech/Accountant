import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

bom_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units).xlsx'
wb = openpyxl.load_workbook(bom_file, data_only=True)

thcp_map = {
    'HRM30A6001': 'DX1TNJ0023',
    'HRS20A6003': 'DX1TNJ0024',
    'DHRM38A6001': 'DX1TNJ0031',
    'SCR18A7501_MC_FSM': 'DX1TNJ0033',
    'GRU30A05-OF-L': 'DX1TNJ0039',
    'GRU30A05-OF-R': 'DX1TNJ0040',
}

ws_main = wb['Đơn vị chính']
ws_sub = wb['đơn vị phụ']

misa_exports = {}
for r in range(6, ws_main.max_row):
    thcp = ws_main.cell(r, 21).value
    code = str(ws_main.cell(r, 3).value or '').strip()
    name = str(ws_main.cell(r, 4).value or '').strip()
    dvt = str(ws_main.cell(r, 10).value or '').strip()
    xuat_sl = ws_main.cell(r, 15).value or 0
    if thcp and code and xuat_sl > 0:
        if thcp not in misa_exports:
            misa_exports[thcp] = {}
        if code not in misa_exports[thcp]:
            misa_exports[thcp][code] = {'name': name, 'main_dvt': dvt, 'main_qty': 0, 'sub_dvt': None, 'sub_qty': 0}
        misa_exports[thcp][code]['main_qty'] += xuat_sl

for r in range(6, ws_sub.max_row):
    thcp = ws_sub.cell(r, 21).value
    code = str(ws_sub.cell(r, 3).value or '').strip()
    dvt = str(ws_sub.cell(r, 10).value or '').strip()
    xuat_sl = ws_sub.cell(r, 15).value or 0
    if thcp and code and xuat_sl > 0:
        if thcp in misa_exports and code in misa_exports[thcp]:
            misa_exports[thcp][code]['sub_dvt'] = dvt
            misa_exports[thcp][code]['sub_qty'] += xuat_sl

for sname, thcp in thcp_map.items():
    ws = wb[sname]
    start_r = None
    for r in range(1, 200):
        if '1.2' in str(ws.cell(r, 1).value or ''):
            start_r = r
            break
            
    bom_codes = set()
    for r in range(start_r + 2, ws.max_row + 1):
        v_a = ws.cell(r, 1).value
        v_c = ws.cell(r, 3).value # name
        v_d = ws.cell(r, 4).value # code
        if v_a and any(k in str(v_a).upper() for k in ['PHẦN 2', '2.', 'TỔNG CỘNG']):
            break
        if v_d:
            c = str(v_d).strip()
            bom_codes.add(c)
            # Also add if numeric e.g. 450 -> CX1NOA0450
            if c.isdigit() and len(c) == 3:
                bom_codes.add(f'CX1NOA0{c}')
                
    misa_map = misa_exports.get(thcp, {})
    unused_misa = set(misa_map.keys()) - bom_codes
    
    print(f'=== {sname} ({thcp}): {len(unused_misa)} MISA codes not in BOM ===')
    for mc in sorted(list(unused_misa)):
        mv = misa_map[mc]
        print(f'   {mc:15s} | Main: {mv["main_qty"]:>8} {mv["main_dvt"]:6s} | Sub: {mv["sub_qty"]:>8} {str(mv["sub_dvt"]):6s} | {mv["name"]}')
    print()
