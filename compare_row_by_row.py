import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

bom_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units).xlsx'
wb = openpyxl.load_workbook(bom_file, data_only=True)

thcp_map = {
    'HRM30A6001': 'DX1TNJ0023',
    'HRS20A6003': 'DX1TNJ0024',
    'DHRM38A6001': 'DX1TNJ0031',
    'SCR18A7501_MC_FSM': 'DX1TNJ0033',
    'GRU30A05-OF-L': 'DX1TNJ0039',
    'GRU30A05-OF-R': 'DX1TNJ0040',
}

ws_main = wb['Đơn vị chính']
ws_sub = wb['đơn vị phụ']

# Build MISA export lookup: thcp -> item_code -> {'name': ..., 'main_qty': ..., 'main_dvt': ..., 'sub_qty': ..., 'sub_dvt': ...}
misa_exports = {}
for r in range(6, ws_main.max_row):
    thcp = ws_main.cell(r, 21).value
    code = str(ws_main.cell(r, 3).value or '').strip()
    name = str(ws_main.cell(r, 4).value or '').strip()
    dvt = str(ws_main.cell(r, 10).value or '').strip()
    xuat_sl = ws_main.cell(r, 15).value or 0
    if thcp and code and xuat_sl > 0:
        if thcp not in misa_exports:
            misa_exports[thcp] = {}
        if code not in misa_exports[thcp]:
            misa_exports[thcp][code] = {'name': name, 'main_dvt': dvt, 'main_qty': 0, 'sub_dvt': None, 'sub_qty': 0}
        misa_exports[thcp][code]['main_qty'] += xuat_sl

for r in range(6, ws_sub.max_row):
    thcp = ws_sub.cell(r, 21).value
    code = str(ws_sub.cell(r, 3).value or '').strip()
    dvt = str(ws_sub.cell(r, 10).value or '').strip()
    xuat_sl = ws_sub.cell(r, 15).value or 0
    if thcp and code and xuat_sl > 0:
        if thcp in misa_exports and code in misa_exports[thcp]:
            misa_exports[thcp][code]['sub_dvt'] = dvt
            misa_exports[thcp][code]['sub_qty'] += xuat_sl

for sname, thcp in thcp_map.items():
    ws = wb[sname]
    start_r = None
    for r in range(1, 200):
        if '1.2' in str(ws.cell(r, 1).value or ''):
            start_r = r
            break
            
    header_r = start_r + 1
    misa_map = misa_exports.get(thcp, {})
    
    print(f'================ {sname} ({thcp}) ================')
    for r in range(header_r + 1, ws.max_row + 1):
        v_a = ws.cell(r, 1).value
        v_b = ws.cell(r, 2).value # STT
        v_c = ws.cell(r, 3).value # Name
        v_d = ws.cell(r, 4).value # Code
        v_f = ws.cell(r, 6).value # Unit
        v_k = ws.cell(r, 11).value # Qty Proj
        v_l = ws.cell(r, 12).value # File XK
        v_m = ws.cell(r, 13).value # Note
        
        if v_a and any(k in str(v_a).upper() for k in ['PHẦN 2', '2.', 'TỔNG CỘNG']):
            break
        if not (v_c or v_d):
            continue
            
        c_code = str(v_d).strip() if v_d else ''
        c_name = str(v_c).strip() if v_c else ''
        c_dvt = str(v_f).strip() if v_f else ''
        
        m_info = misa_map.get(c_code)
        m_str = ''
        if m_info:
            m_str = f'MISA direct: Main={m_info["main_qty"]} {m_info["main_dvt"]}, Sub={m_info["sub_qty"]} {m_info["sub_dvt"]}'
        else:
            m_str = 'MISA direct: NOT FOUND'
            
        print(f'R{r:3d} | STT:{str(v_b):>3} | Code:{c_code:12s} | DVT:{c_dvt:5s} | Định mức:{str(v_k):>8} | File XK:{str(v_l):>8} | {m_str} | Name:{c_name[:25]} | Note:{str(v_m) if v_m else ""}')
