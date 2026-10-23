import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

backup_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units)_backup.xlsx'

# Load both: formula wb for writing (preserving formulas elsewhere), data wb for reading evaluated values
wb = openpyxl.load_workbook(backup_file, data_only=False)
wb_data = openpyxl.load_workbook(backup_file, data_only=True)

thcp_map = {
    'HRM30A6001': 'DX1TNJ0023',
    'HRS20A6003': 'DX1TNJ0024',
    'DHRM38A6001': 'DX1TNJ0031',
    'SCR18A7501_MC_FSM': 'DX1TNJ0033',
    'GRU30A05-OF-L': 'DX1TNJ0039',
    'GRU30A05-OF-R': 'DX1TNJ0040',
}

ws_main = wb_data['Đơn vị chính']
ws_sub = wb_data['đơn vị phụ']

# Build comprehensive lookup dictionaries
xuat_chinh_thcp = {}
xuat_tieuhao_thcp = {}
xuat_tieuhao_nothcp = {}
nhap_thua_thcp = {}
nhap_thua_nothcp = {}

# Process ws_main
for r in range(6, ws_main.max_row):
    code = str(ws_main.cell(r, 3).value or '').strip()
    name = str(ws_main.cell(r, 4).value or '').strip()
    so_ct = str(ws_main.cell(r, 8).value or '').strip()
    dien_giai = str(ws_main.cell(r, 9).value or '').strip()
    dvt = str(ws_main.cell(r, 10).value or '').strip()
    sl_nhap = ws_main.cell(r, 13).value or 0
    sl_xuat = ws_main.cell(r, 15).value or 0
    tk_du = str(ws_main.cell(r, 20).value or '').strip()
    thcp = str(ws_main.cell(r, 21).value or '').strip()
    kmcp = str(ws_main.cell(r, 23).value or '').strip()
    
    if not code:
        continue
        
    # Check Nhập kho thừa từ sản xuất (TK 154 hoặc Diễn giải ghi rõ NVL thừa / thừa khác)
    if sl_nhap > 0:
        is_thua = (tk_du == '154' or 'nvl thừa' in dien_giai.lower() or 'thừa khác' in dien_giai.lower() or ('thừa' in dien_giai.lower() and 'giao thừa' not in dien_giai.lower()))
        if is_thua:
            if thcp:
                key = (thcp, code)
                if key not in nhap_thua_thcp:
                    nhap_thua_thcp[key] = {'name': name, 'main_dvt': dvt, 'main_qty': 0, 'sub_dvt': None, 'sub_qty': 0, 'rows': []}
                nhap_thua_thcp[key]['main_qty'] += sl_nhap
                nhap_thua_thcp[key]['rows'].append((r, sl_nhap, dvt, dien_giai, so_ct))
            else:
                if code not in nhap_thua_nothcp:
                    nhap_thua_nothcp[code] = {'name': name, 'main_dvt': dvt, 'main_qty': 0, 'sub_dvt': None, 'sub_qty': 0, 'rows': []}
                nhap_thua_nothcp[code]['main_qty'] += sl_nhap
                nhap_thua_nothcp[code]['rows'].append((r, sl_nhap, dvt, dien_giai, so_ct))
                
    # Check Xuất kho
    if sl_xuat > 0:
        is_tieu_hao = ('tiêu hao' in dien_giai.lower() or 'tieu hao' in dien_giai.lower() or kmcp.startswith('SXC'))
        if thcp:
            key = (thcp, code)
            if is_tieu_hao:
                if key not in xuat_tieuhao_thcp:
                    xuat_tieuhao_thcp[key] = {'name': name, 'main_dvt': dvt, 'main_qty': 0, 'sub_dvt': None, 'sub_qty': 0, 'rows': []}
                xuat_tieuhao_thcp[key]['main_qty'] += sl_xuat
                xuat_tieuhao_thcp[key]['rows'].append((r, sl_xuat, dvt, dien_giai, so_ct))
            else:
                if key not in xuat_chinh_thcp:
                    xuat_chinh_thcp[key] = {'name': name, 'main_dvt': dvt, 'main_qty': 0, 'sub_dvt': None, 'sub_qty': 0, 'rows': []}
                xuat_chinh_thcp[key]['main_qty'] += sl_xuat
                xuat_chinh_thcp[key]['rows'].append((r, sl_xuat, dvt, dien_giai, so_ct))
        else:
            if is_tieu_hao:
                if code not in xuat_tieuhao_nothcp:
                    xuat_tieuhao_nothcp[code] = {'name': name, 'main_dvt': dvt, 'main_qty': 0, 'sub_dvt': None, 'sub_qty': 0, 'rows': []}
                xuat_tieuhao_nothcp[code]['main_qty'] += sl_xuat
                xuat_tieuhao_nothcp[code]['rows'].append((r, sl_xuat, dvt, dien_giai, so_ct))

# Process ws_sub for sub_qty and sub_dvt
for r in range(6, ws_sub.max_row):
    code = str(ws_sub.cell(r, 3).value or '').strip()
    dvt = str(ws_sub.cell(r, 10).value or '').strip()
    sl_nhap = ws_sub.cell(r, 13).value or 0
    sl_xuat = ws_sub.cell(r, 15).value or 0
    dien_giai = str(ws_sub.cell(r, 9).value or '').strip()
    tk_du = str(ws_sub.cell(r, 20).value or '').strip()
    thcp = str(ws_sub.cell(r, 21).value or '').strip()
    kmcp = str(ws_sub.cell(r, 23).value or '').strip()
    
    if not code:
        continue
        
    if sl_nhap > 0:
        is_thua = (tk_du == '154' or 'nvl thừa' in dien_giai.lower() or 'thừa khác' in dien_giai.lower() or ('thừa' in dien_giai.lower() and 'giao thừa' not in dien_giai.lower()))
        if is_thua:
            if thcp and (thcp, code) in nhap_thua_thcp:
                nhap_thua_thcp[(thcp, code)]['sub_dvt'] = dvt
                nhap_thua_thcp[(thcp, code)]['sub_qty'] += sl_nhap
            elif not thcp and code in nhap_thua_nothcp:
                nhap_thua_nothcp[code]['sub_dvt'] = dvt
                nhap_thua_nothcp[code]['sub_qty'] += sl_nhap
                
    if sl_xuat > 0:
        is_tieu_hao = ('tiêu hao' in dien_giai.lower() or 'tieu hao' in dien_giai.lower() or kmcp.startswith('SXC'))
        if thcp:
            key = (thcp, code)
            if is_tieu_hao and key in xuat_tieuhao_thcp:
                xuat_tieuhao_thcp[key]['sub_dvt'] = dvt
                xuat_tieuhao_thcp[key]['sub_qty'] += sl_xuat
            elif not is_tieu_hao and key in xuat_chinh_thcp:
                xuat_chinh_thcp[key]['sub_dvt'] = dvt
                xuat_chinh_thcp[key]['sub_qty'] += sl_xuat
        else:
            if is_tieu_hao and code in xuat_tieuhao_nothcp:
                xuat_tieuhao_nothcp[code]['sub_dvt'] = dvt
                xuat_tieuhao_nothcp[code]['sub_qty'] += sl_xuat

# Master substitutions and mappings
substitutions = {
    'CX1TBM0392': ['CX1HTP0610'], # Mặt bích tròn DN80
    'CX1TBM0393': ['CX1HTP0611'], # Mặt bích vuông DN80
    'CX1TBM0142': ['CX1HTP0548'], # Mặt bích tròn DN50
    'CX1TBM0365': ['CX1HTP0609'], # Mặt bích vuông DN50
    'CX1TBM0549': ['CX1HTP0549'], # Mặt bích tròn DN25
    'CX1TBM0141': ['CX1HTP0549'], # Mặt bích tròn DN25
    'CX1PLO0441': ['CX1HTP0549'], # Mặt bích DN25 Inox 125x125x5mm
    'CX1KAN0158': ['CX1MAX0691'], # Xích 35-1R (merge KANA + Maxton)
    'CX1KAN0157': ['CX1NOA0470', 'CX1STA0318', 'CX1DID0168'], # Xích 40-1R (merge KANA + Maxton + STAR + DID)
    'CX1NOA0470': ['CX1STA0318'], # Xích 40-1R
    'CX1NOA0300': ['CX1NOA0681'], # Long đền M14x30 (1mm + 2mm)
    'CX1NOA0277': ['CX1NOA0682'], # Long đền M16x30 (1mm + 2mm)
    'CX1NOA0455': ['CX1NOA0680'], # Long đền M8x19 (1mm + 2mm)
    'CX1NOA0641': ['CX1NOA0680'], # Long đền M8x19x1 (SCR)
    'CX1NOA0149': ['CX1NOA0537'], # Nhựa POM DN40 (merge)
    'CX1NOA0378': ['CX1KLO0677', 'CX1NOA0377'], # Măng xông DN80 / Ron Teflon DN80
    'CX1NTN0250': ['CX1PBC0429'], # Bạc đạn 6808 (NTN -> PBC)
    'CX1SKF0111': ['CX1UKB0692'], # Bạc đạn 6206 (SKF -> UKB)
    'CX1NOA0130': ['CX1MAX0690'], # Nhông xích 40B-12T
    'CX1KLO0254': ['CX1KLO0600'], # Nắp hộp giảm tốc
    'CX1KLO0422': ['CX1KLO0596'], # MC Khâu 2
    'CX1VVU0584': ['CX1VVU0585'], # Pat điều khiển bộ giảm tốc
    'CX1NOA0468': ['CX1NOA0704'], # Bulong M8x15 / M8x16
    'CX1NOA0640': ['CX1NOA0450'], # Bulong M5x15 / M5x16
    '450': ['CX1NOA0450'],
    '700': ['CX1NOA0700'],
    '702': ['CX1NOA0702'],
    '703': ['CX1NOA0703'],
    '704': ['CX1NOA0704'],
    '309': ['CX1NOA0309'],
    'CX1KLO0605': ['CX1KLO0605'],
    'CX1NOA0635': ['CX1KLO0605'], # Ổ đỡ bạc đạn 6002
    'CX1RBI0445': ['CX1FBJ0705'], # Bạc đạn 696 ZZ
    'CX1DPI0445': ['CX1FBJ0705', 'CX1RBI0564'],
    'CX1KLO0413': ['CX1HTP0628', 'CX1HTP0629'], # Lưỡi gài / chốt giữ
    'CX1HTP0491': ['CX1HTP0628'], # LƯỠI GÀI 109.8*22.6*21.6mm
    'CX1NOA0105': ['CX1NOA0013'], # Cây ty tròn đặc 14mm / 18mm
    'CX1NOA0546': ['CX1NOA0013'], # Cây ty tròn đặc 14mm
    'CX1KLO0420': ['CX1SFM0683'], # Khớp xoay SVE12S02
    'CX1PLY0447': ['CX1PLY0698'], # Lò xo cài cóc
    'CX1NOA0650': ['CX1PLY0698'], # Lò xo cài cóc
    'CX1DWH0503': ['CX1DWH0502'], # Lò xo xoắn cuộn
    'CX1TBM0528': ['CX1VVU0571'], # Hộp lò xo P2 phi 486x31x3mm
    'CX1LKB0693': ['CX1LKB0694'], # Gối đỡ UCP212
    'CX1CUX0029': ['CX1LKB0694', 'CX1NOA0059'], # Gối đỡ UCP212 CUXB / Co hàn 45
    'CX1SKF0112': ['CX1PBC0494', 'CX1RBI0562'], # Vòng bi 6006-2RS
    'CX1F&D0295': ['CX1PBC0494', 'CX1RBI0562'], # Vòng bi 6006-2RS
    'CX1NOA0172': ['CX1NOA0369'], # Khớp quay cuốn cáp
    'CX1NOA0159': ['CX1NOA0686'], # Bulong đầu dù bake M6x10
    'CX1JTE0511': ['CX1RBI0564'], # Vòng bi 626
    'CX1KLO0421': ['CX1HTP0629'], # Chốt giữ lưỡi gài
    'CX1TBM0534': ['CX1VVU0583'], # Khung định hướng ống
    'CX1NOA0232': ['CX1NOA0008'], # Bulong M10x25 (xuất thực tế 10x20)
    'CX1NOA0060': ['CX1NOA0160'], # Co hàn 90 DN50 SCH40
    'CX1ATL0395': ['CX1MNA0023'], # Que hàn Atlantic -> Menam
    '53': ['CX1NOA0053'], # Bulong M10x40
    '629': ['CX1HTP0629'], # Chốt giữ lưỡi gài
    '683': ['CX1SFM0683'], # Khớp xoay
    '583': ['CX1VVU0583'], # Khung định hướng ống
    '698': ['CX1PLY0698'], # Lò xo bật cài cóc
    '502': ['CX1DWH0502'], # Lò xo xoắn cuộn
    '440': ['CX1NOA0440'], # Bulong đầu bằng 8x70
    '517': ['CX1NOA0517'], # Bulong 6x80
    '679': ['CX1NOA0679'],
    'CX1NOA0380': ['CX1NOA0160'], # Co hàn 90 DN50
    'CX1NOA0148': ['CX1NOA0281'], # Bulong M6x6
    'CX1NOA0643': ['CX1NOA0703', 'CX1NOA0700'], # Bulong M5x16 / Đinh tán M5x16 thay M5x25
}

name_substitutions = {
    'que hàn inox er308l': 'CX1MNA0049',
    'que hàn inox atlantic': 'CX1ATL0395',
    'chg-308': 'CX1ATL0395',
    'khóa xích 35': 'CX1NOA0272',
    'khóa xích 40': 'CX1NOA0273',
    'đinh rút 4mm': 'CX1NOA0341',
    'bulong đầu dù vit bake': 'CX1NOA0686',
    'bulong đầu dẹp 5x16': 'CX1NOA0703',
}

inferred_dvt_map = {
    ('DHRM38A6001', '440'): 'Cái',
    ('DHRM38A6001', '53'): 'Cái',
    ('DHRM38A6001', '629'): 'Cái',
    ('DHRM38A6001', '683'): 'Cái',
    ('DHRM38A6001', '583'): 'Cái',
    ('DHRM38A6001', '698'): 'Cái',
    ('DHRM38A6001', '502'): 'Cái',
    ('SCR18A7501_MC_FSM', '517'): 'Cái',
    ('SCR18A7501_MC_FSM', 'CX1NOA0638'): 'Cái',
    ('SCR18A7501_MC_FSM', 'CX1KLO0608'): 'Cái',
    ('SCR18A7501_MC_FSM', '450'): 'Cái',
    ('GRU30A05-OF-L', '700'): 'Cái',
    ('GRU30A05-OF-L', '704'): 'Cái',
    ('GRU30A05-OF-L', '702'): 'Cái',
    ('GRU30A05-OF-R', '450'): 'Cái',
    ('GRU30A05-OF-R', '703'): 'Cái',
    ('GRU30A05-OF-R', '704'): 'Cái',
    ('GRU30A05-OF-R', '309'): 'Cái',
}

def convert_misa_qty(m_obj, bom_dvt):
    bd = bom_dvt.lower().strip()
    m_main_dvt = m_obj['main_dvt'].lower().strip()
    m_sub_dvt = (m_obj['sub_dvt'] or '').lower().strip()
    
    # Case 1: Length in mm (BOM is mm, MISA is m/mét)
    if bd in ['mm']:
        if m_main_dvt in ['m', 'mét']:
            return m_obj['main_qty'] * 1000
        elif m_sub_dvt in ['m', 'mét']:
            return m_obj['sub_qty'] * 1000
        elif m_sub_dvt in ['hộp'] and m_main_dvt in ['m', 'mét']:
            return m_obj['main_qty'] * 1000
        elif m_main_dvt in ['cây'] and m_sub_dvt in ['m', 'mét']:
            return m_obj['sub_qty'] * 1000
            
    # Case 2: Length in m
    if bd in ['m', 'mét']:
        if m_main_dvt in ['m', 'mét']:
            return m_obj['main_qty']
        elif m_sub_dvt in ['m', 'mét']:
            return m_obj['sub_qty']
            
    # Case 3: Welding rods in que / cây (BOM is cây/que, MISA main is kg, sub is que)
    if bd in ['que', 'cây', 'cay']:
        if m_sub_dvt in ['que', 'cây', 'cay']:
            return m_obj['sub_qty']
        elif m_main_dvt in ['que', 'cây', 'cay']:
            return m_obj['main_qty']
            
    # Case 4: Pieces / Items
    if bd in ['cái', 'cái', 'con', 'bộ', 'chai', 'lon', 'hộp', 'tấm', 'đôi', 'mũi', 'cặp']:
        if m_main_dvt in ['cái', 'cái', 'con', 'bộ', 'chai', 'lon', 'hộp', 'tấm', 'đôi', 'mũi', 'cặp']:
            return m_obj['main_qty']
        elif m_sub_dvt in ['cái', 'cái', 'con', 'bộ', 'chai', 'lon', 'hộp', 'tấm', 'đôi', 'mũi', 'cặp']:
            return m_obj['sub_qty']
        else:
            return m_obj['main_qty']
            
    return m_obj['main_qty']

# Styles
font_family = 'Times New Roman'
header_font = Font(name=font_family, size=10, bold=True)
title_font = Font(name=font_family, size=13, bold=True)
sub_font = Font(name=font_family, size=10, italic=True)
regular_font = Font(name=font_family, size=10)
bold_font = Font(name=font_family, size=10, bold=True)
italic_font = Font(name=font_family, size=10, italic=True)

header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid') # Soft blue
match_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid') # Soft green (Khớp / Nằm trong ĐM)
over_fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')  # Soft orange (Xuất thừa)
under_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid') # Soft yellow (Xuất thiếu nhiều)
zero_fill = PatternFill(start_color='EDEDED', end_color='EDEDED', fill_type='solid')  # Gray (Chưa xuất / Không xuất)
extra_fill = PatternFill(start_color='E6B8B7', end_color='E6B8B7', fill_type='solid') # Soft reddish (Ngoài BOM)

thin_border_side = Side(border_style='thin', color='B0B0B0')
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
double_bottom_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style='double', color='000000'))

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')

num_fmt_qty = '#,##0.00_);(#,##0.00);"-";@'
num_fmt_int = '#,##0_);(#,##0);"-";@'
num_fmt_pct = '0.0%'

all_master_rows = []

for sname, thcp in thcp_map.items():
    ws = wb[sname]
    ws_d = wb_data[sname]
    
    start_r = None
    for r in range(1, 200):
        if '1.2' in str(ws_d.cell(r, 1).value or ''):
            start_r = r
            break
            
    header_r = start_r + 1
    
    # Safely unmerge any cells in cols 12-22 for rows >= header_r
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= header_r and m.max_col >= 12 and m.min_col <= 22:
            ws.unmerge_cells(str(m))
            
    # Product sheet headers
    prod_headers = [
        (12, 'Xuất NVL chính (Theo TP)', 20),
        (13, 'Xuất tiêu hao theo TP', 20),
        (14, 'Tiêu hao chung (Không gắn mã TP)', 24),
        (15, 'Nhập lại kho (NVL thừa)', 20),
        (16, 'Tổng thực xuất Net (Xuất - Nhập thừa)', 24),
        (17, 'Chênh lệch (Net - Định mức)', 22),
        (18, 'Mã MISA thực xuất / Ghi chú đối chiếu & Quy đổi', 48),
    ]
    
    for c_idx, h_title, h_width in prod_headers:
        c = ws.cell(header_r, c_idx, h_title)
        c.font = header_font
        c.fill = header_fill
        c.alignment = align_center
        c.border = thin_border
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = h_width
        
    claimed_chinh = set()
    claimed_th_tp = set()
    claimed_nhap = set()
    
    for r in range(header_r + 1, header_r + 150):
        name = ws_d.cell(r, 3).value
        code = ws_d.cell(r, 4).value
        dvt = ws_d.cell(r, 6).value
        qty_proj = ws_d.cell(r, 11).value
        note = ws_d.cell(r, 13).value
        stt_bom = ws_d.cell(r, 2).value
        
        if not (name and str(name).strip()):
            continue
            
        c_name = str(name).strip()
        c_code = str(code).strip() if code and str(code).strip() != '#N/A' else ''
        c_dvt = str(dvt).strip() if dvt else ''
        
        # Point 1: Unify unit according to BOM
        if not c_dvt:
            if (sname, c_code) in inferred_dvt_map:
                c_dvt = inferred_dvt_map[(sname, c_code)]
            elif any(k in c_name.lower() for k in ['bulong', 'tán', 'ốc', 'tay quay', 'long đền', 'đinh tán', 'chốt', 'phe gài', 'lò xo', 'khớp xoay', 'khung']):
                c_dvt = 'Cái'
            elif 'que hàn' in c_name.lower():
                c_dvt = 'Cây'
            elif any(k in c_name.lower() for k in ['ống', 'ty', 'nhôm', 'thanh tròn', 'xích']):
                c_dvt = 'mm'
            else:
                c_dvt = 'Cái'
                
            # Fill back unified unit to BOM product sheet Column F
            cell_f = ws.cell(r, 6, c_dvt)
            cell_f.alignment = align_center
            cell_f.font = regular_font
            
        c_proj = 0.0
        if qty_proj is not None:
            try:
                c_proj = float(qty_proj)
            except:
                c_proj = 0.0
                
        candidates = []
        if c_code:
            candidates.append(c_code)
            if c_code in substitutions:
                candidates.extend(substitutions[c_code])
            if c_code.isdigit() and len(c_code) == 3:
                candidates.append(f'CX1NOA0{c_code}')
        else:
            for k_name, m_code in name_substitutions.items():
                if k_name in c_name.lower():
                    candidates.append(m_code)
                    if m_code in substitutions:
                        candidates.extend(substitutions[m_code])
                        
        row_chinh = 0.0
        row_th_tp = 0.0
        row_th_chung = 0.0
        row_nhap = 0.0
        
        matched_misa_desc = []
        
        for cand in candidates:
            # 1. Xuất NVL chính (Theo TP)
            if (thcp, cand) in xuat_chinh_thcp and cand not in claimed_chinh:
                m_obj = xuat_chinh_thcp[(thcp, cand)]
                c_qty = convert_misa_qty(m_obj, c_dvt)
                row_chinh += c_qty
                claimed_chinh.add(cand)
                matched_misa_desc.append(f'Chính: {cand} ({c_qty:g} {c_dvt})')
                
            # 2. Xuất tiêu hao theo TP
            if (thcp, cand) in xuat_tieuhao_thcp and cand not in claimed_th_tp:
                m_obj = xuat_tieuhao_thcp[(thcp, cand)]
                c_qty = convert_misa_qty(m_obj, c_dvt)
                row_th_tp += c_qty
                claimed_th_tp.add(cand)
                matched_misa_desc.append(f'Tiêu hao TP: {cand} ({c_qty:g} {c_dvt})')
                
            # 3. Tiêu hao chung (Không gắn mã TP)
            if cand in xuat_tieuhao_nothcp:
                m_obj = xuat_tieuhao_nothcp[cand]
                c_qty = convert_misa_qty(m_obj, c_dvt)
                row_th_chung += c_qty
                
            # 4. Nhập lại kho (NVL thừa)
            if (thcp, cand) in nhap_thua_thcp and cand not in claimed_nhap:
                m_obj = nhap_thua_thcp[(thcp, cand)]
                c_qty = convert_misa_qty(m_obj, c_dvt)
                row_nhap += c_qty
                claimed_nhap.add(cand)
                matched_misa_desc.append(f'Nhập thừa: {cand} ({c_qty:g} {c_dvt})')
                
        # Fill cells in Product sheet
        # Col L: Xuất chính
        cell_l = ws.cell(r, 12, row_chinh)
        cell_l.font = bold_font
        cell_l.alignment = align_right
        cell_l.number_format = num_fmt_qty if isinstance(row_chinh, float) and not row_chinh.is_integer() else num_fmt_int
        cell_l.border = thin_border
        
        # Col M: Xuất tiêu hao theo TP
        cell_m = ws.cell(r, 13, row_th_tp)
        cell_m.font = bold_font if row_th_tp > 0 else regular_font
        cell_m.alignment = align_right
        cell_m.number_format = num_fmt_qty if isinstance(row_th_tp, float) and not row_th_tp.is_integer() else num_fmt_int
        cell_m.border = thin_border
        
        # Col N: Tiêu hao chung (Không gắn mã TP)
        cell_n = ws.cell(r, 14, row_th_chung)
        cell_n.font = bold_font if row_th_chung > 0 else regular_font
        cell_n.alignment = align_right
        cell_n.number_format = num_fmt_qty if isinstance(row_th_chung, float) and not row_th_chung.is_integer() else num_fmt_int
        cell_n.border = thin_border
        
        # Col O: Nhập lại kho (NVL thừa)
        cell_o = ws.cell(r, 15, row_nhap)
        cell_o.font = bold_font if row_nhap > 0 else regular_font
        cell_o.alignment = align_right
        cell_o.number_format = num_fmt_qty if isinstance(row_nhap, float) and not row_nhap.is_integer() else num_fmt_int
        cell_o.border = thin_border
        
        # Col P: Tổng thực xuất Net = (L{r}+M{r})-O{r}
        cell_p = ws.cell(r, 16)
        cell_p.value = f'=(L{r}+M{r})-O{r}'
        cell_p.font = bold_font
        cell_p.alignment = align_right
        cell_p.number_format = num_fmt_qty
        cell_p.border = thin_border
        
        # Col Q: Chênh lệch = P{r}-K{r}
        cell_q = ws.cell(r, 17)
        cell_q.value = f'=P{r}-K{r}'
        cell_q.font = bold_font
        cell_q.alignment = align_right
        cell_q.number_format = num_fmt_qty
        cell_q.border = thin_border
        
        net_val = (row_chinh + row_th_tp) - row_nhap
        diff_val = net_val - c_proj
        
        # Point 2: Rule "Những cái nào xuất dưới định mức không nhiều thì nằm trong định mức và không phải xuất thiếu"
        if c_proj > 0:
            if abs(diff_val) < 0.001:
                status_str = 'Khớp định mức'
                cell_fill = match_fill
            elif diff_val > 0.001:
                status_str = f'Xuất thừa (+{diff_val:g})'
                cell_fill = over_fill
            else:
                # Net < c_proj
                if net_val == 0:
                    status_str = 'Chưa xuất kho'
                    cell_fill = zero_fill
                elif net_val >= 0.85 * c_proj:
                    # Xuất dưới định mức không nhiều (trong ngưỡng dung sai 5-15%): Nằm trong định mức!
                    status_str = f'Nằm trong định mức ({diff_val:g})'
                    cell_fill = match_fill # Soft Green như Khớp định mức!
                else:
                    status_str = f'Xuất thiếu ({diff_val:g})'
                    cell_fill = under_fill # Soft Yellow
        else:
            if net_val > 0:
                status_str = f'Xuất ngoài ĐM (+{net_val:g})'
                cell_fill = extra_fill
            else:
                status_str = 'Không xuất (ĐM = 0)'
                cell_fill = zero_fill
                
        cell_q.fill = cell_fill
        cell_p.fill = cell_fill
        
        # Notes
        note_parts = []
        if matched_misa_desc:
            note_parts.append('MISA: ' + '; '.join(matched_misa_desc))
        else:
            note_parts.append('MISA: Chưa xuất riêng cho TP này')
            
        if row_th_chung > 0:
            note_parts.append(f'Tiêu hao chung xưởng: {row_th_chung:g} {c_dvt}')
            
        # Check unassigned excess returns
        unassigned_ret = 0.0
        for cand in candidates:
            if cand in nhap_thua_nothcp:
                unassigned_ret += convert_misa_qty(nhap_thua_nothcp[cand], c_dvt)
        if unassigned_ret > 0:
            note_parts.append(f'Lưu ý: Có {unassigned_ret:g} {c_dvt} nhập thừa chung chưa gán mã TP trên MISA')
            
        if note and str(note).strip():
            note_parts.append(f'Ghi chú file: {str(note).strip()}')
            
        full_note = ' | '.join(note_parts)
        cell_r = ws.cell(r, 18, full_note)
        cell_r.font = regular_font
        cell_r.alignment = align_left
        cell_r.border = thin_border
        
        # Collect for Master sheet
        all_master_rows.append({
            'product': sname,
            'thcp': thcp,
            'stt_bom': stt_bom,
            'bom_code': c_code if c_code else '(Chưa có mã)',
            'bom_name': c_name,
            'bom_dvt': c_dvt,
            'qty_proj': c_proj,
            'misa_codes': ', '.join([c for c in candidates if (thcp, c) in xuat_chinh_thcp or (thcp, c) in xuat_tieuhao_thcp or (thcp, c) in nhap_thua_thcp]) or '(Chưa xuất)',
            'misa_names': c_name,
            'misa_dvt': c_dvt,
            'xk_chinh': row_chinh,
            'xk_th_tp': row_th_tp,
            'th_chung': row_th_chung,
            'nhap_thua': row_nhap,
            'net_xuat': net_val,
            'diff': diff_val,
            'status': status_str,
            'cell_fill': cell_fill,
            'note': full_note
        })
        
    # Append any remaining unclaimed MISA items for this product
    unclaimed_chinh = [c for (t, c) in xuat_chinh_thcp if t == thcp and c not in claimed_chinh]
    unclaimed_th_tp = [c for (t, c) in xuat_tieuhao_thcp if t == thcp and c not in claimed_th_tp]
    all_unclaimed = set(unclaimed_chinh + unclaimed_th_tp)
    
    for uc in all_unclaimed:
        u_chinh = 0.0
        u_th_tp = 0.0
        u_nhap = 0.0
        u_name = ''
        u_dvt = ''
        
        if (thcp, uc) in xuat_chinh_thcp:
            m_obj = xuat_chinh_thcp[(thcp, uc)]
            u_chinh = m_obj['main_qty']
            u_name = m_obj['name']
            u_dvt = m_obj['main_dvt']
        if (thcp, uc) in xuat_tieuhao_thcp:
            m_obj = xuat_tieuhao_thcp[(thcp, uc)]
            u_th_tp = m_obj['main_qty']
            if not u_name:
                u_name = m_obj['name']
                u_dvt = m_obj['main_dvt']
        if (thcp, uc) in nhap_thua_thcp:
            m_obj = nhap_thua_thcp[(thcp, uc)]
            u_nhap = m_obj['main_qty']
            
        u_net = (u_chinh + u_th_tp) - u_nhap
        all_master_rows.append({
            'product': sname,
            'thcp': thcp,
            'stt_bom': '(Ngoài BOM)',
            'bom_code': '(Không có)',
            'bom_name': f'(Ngoài BOM) {u_name}',
            'bom_dvt': u_dvt,
            'qty_proj': 0.0,
            'misa_codes': uc,
            'misa_names': u_name,
            'misa_dvt': u_dvt,
            'xk_chinh': u_chinh,
            'xk_th_tp': u_th_tp,
            'th_chung': 0.0,
            'nhap_thua': u_nhap,
            'net_xuat': u_net,
            'diff': u_net,
            'status': f'Xuất ngoài ĐM (+{u_net:g})',
            'cell_fill': extra_fill,
            'note': f'MISA thực xuất: {uc} ({u_name[:30]}, xuất {u_net:g} {u_dvt}) phát sinh thực tế tại phân xưởng, không có trong danh mục BOM kỹ thuật'
        })
        
    print(f'Sheet {sname} updated: Unified BOM DVT & status rule applied.')

# ==========================================
# MASTER SHEET: DOI_CHIEU_BOM_MISA
# ==========================================
master_sheet_name = 'DOI_CHIEU_BOM_MISA'
if master_sheet_name in wb.sheetnames:
    del wb[master_sheet_name]
ws_master = wb.create_sheet(title=master_sheet_name, index=0)

# Title Block
ws_master['A1'] = 'BẢNG TỔNG HỢP ĐỐI CHIẾU ĐỊNH MỨC NGUYÊN VẬT LIỆU (BOM) VS XUẤT KHO THỰC TẾ (MISA)'
ws_master['A1'].font = title_font
ws_master['A2'] = 'Dự án: FSM Master (Gói 2 - 10 Units) | Phạm vi: 6 Thành phẩm chính thức | ĐVT: Thống nhất hoàn toàn theo BOM | Quy chuẩn: Xuất dưới ĐM không nhiều tính là Nằm trong định mức'
ws_master['A2'].font = sub_font

master_headers = [
    ('STT', 6, align_center),
    ('Sản phẩm thành phẩm', 18, align_left),
    ('Mã THCP (MISA)', 14, align_center),
    ('STT BOM', 10, align_center),
    ('Mã vật tư (BOM)', 16, align_left),
    ('Tên vật tư (BOM)', 36, align_left),
    ('ĐVT (Thống nhất theo BOM)', 16, align_center),
    ('Định mức dự án', 15, align_right),
    ('Mã MISA thực xuất', 22, align_left),
    ('Tên vật tư MISA thực xuất', 34, align_left),
    ('ĐVT MISA gốc', 14, align_center),
    ('Xuất NVL chính (Theo TP)', 18, align_right),
    ('Xuất tiêu hao theo TP', 18, align_right),
    ('Tiêu hao chung (Không gắn mã TP)', 22, align_right),
    ('Nhập lại kho (NVL thừa)', 18, align_right),
    ('Tổng thực xuất Net (Xuất - Nhập thừa)', 22, align_right),
    ('Chênh lệch (Net - Định mức)', 18, align_right),
    ('Tỷ lệ đáp ứng', 12, align_right),
    ('Đánh giá trạng thái', 22, align_center),
    ('Ghi chú chi tiết & Quy đổi', 52, align_left),
]

for col_idx, (h_title, h_width, h_align) in enumerate(master_headers, 1):
    c = ws_master.cell(4, col_idx, h_title)
    c.font = header_font
    c.fill = header_fill
    c.alignment = align_center
    c.border = thin_border
    col_letter = get_column_letter(col_idx)
    ws_master.column_dimensions[col_letter].width = h_width

curr_master_row = 5
prev_product = None

for idx, item in enumerate(all_master_rows, 1):
    is_new_product = (item['product'] != prev_product)
    prev_product = item['product']
    r = curr_master_row
    
    # Col A: STT
    c_stt = ws_master.cell(r, 1, idx)
    c_stt.alignment = align_center
    c_stt.font = regular_font
    c_stt.border = thin_border
    
    # Col B: Product
    c_prod = ws_master.cell(r, 2, item['product'])
    c_prod.alignment = align_left
    c_prod.font = bold_font if is_new_product else regular_font
    c_prod.border = thin_border
    
    # Col C: THCP
    c_thcp = ws_master.cell(r, 3, item['thcp'])
    c_thcp.alignment = align_center
    c_thcp.font = bold_font if is_new_product else regular_font
    c_thcp.border = thin_border
    
    # Col D: STT BOM
    c_sb = ws_master.cell(r, 4, item['stt_bom'])
    c_sb.alignment = align_center
    c_sb.font = regular_font
    c_sb.border = thin_border
    
    # Col E: BOM Code
    c_bcode = ws_master.cell(r, 5, item['bom_code'])
    c_bcode.alignment = align_left
    c_bcode.font = bold_font
    c_bcode.border = thin_border
    
    # Col F: BOM Name
    c_bname = ws_master.cell(r, 6, item['bom_name'])
    c_bname.alignment = align_left
    c_bname.font = regular_font
    c_bname.border = thin_border
    
    # Col G: Unified BOM DVT
    c_bdvt = ws_master.cell(r, 7, item['bom_dvt'])
    c_bdvt.alignment = align_center
    c_bdvt.font = bold_font
    c_bdvt.border = thin_border
    
    # Col H: BOM Qty Proj
    c_proj = ws_master.cell(r, 8, item['qty_proj'])
    c_proj.alignment = align_right
    c_proj.font = bold_font
    c_proj.number_format = num_fmt_qty if isinstance(item['qty_proj'], float) and not item['qty_proj'].is_integer() else num_fmt_int
    c_proj.border = thin_border
    
    # Col I: MISA Code
    c_mcode = ws_master.cell(r, 9, item['misa_codes'])
    c_mcode.alignment = align_left
    c_mcode.font = bold_font if item['misa_codes'] != '(Chưa xuất)' else italic_font
    c_mcode.border = thin_border
    
    # Col J: MISA Name
    c_mname = ws_master.cell(r, 10, item['misa_names'])
    c_mname.alignment = align_left
    c_mname.font = regular_font
    c_mname.border = thin_border
    
    # Col K: MISA DVT
    c_mdvt = ws_master.cell(r, 11, item['misa_dvt'])
    c_mdvt.alignment = align_center
    c_mdvt.font = regular_font
    c_mdvt.border = thin_border
    
    # Col L: Xuất NVL chính (Theo TP)
    c_l = ws_master.cell(r, 12, item['xk_chinh'])
    c_l.alignment = align_right
    c_l.font = bold_font if item['xk_chinh'] > 0 else regular_font
    c_l.number_format = num_fmt_qty if isinstance(item['xk_chinh'], float) and not item['xk_chinh'].is_integer() else num_fmt_int
    c_l.border = thin_border
    
    # Col M: Xuất tiêu hao theo TP
    c_m = ws_master.cell(r, 13, item['xk_th_tp'])
    c_m.alignment = align_right
    c_m.font = bold_font if item['xk_th_tp'] > 0 else regular_font
    c_m.number_format = num_fmt_qty if isinstance(item['xk_th_tp'], float) and not item['xk_th_tp'].is_integer() else num_fmt_int
    c_m.border = thin_border
    
    # Col N: Tiêu hao chung (Không gắn mã TP)
    c_n = ws_master.cell(r, 14, item['th_chung'])
    c_n.alignment = align_right
    c_n.font = bold_font if item['th_chung'] > 0 else regular_font
    c_n.number_format = num_fmt_qty if isinstance(item['th_chung'], float) and not item['th_chung'].is_integer() else num_fmt_int
    c_n.border = thin_border
    
    # Col O: Nhập lại kho (NVL thừa)
    c_o = ws_master.cell(r, 15, item['nhap_thua'])
    c_o.alignment = align_right
    c_o.font = bold_font if item['nhap_thua'] > 0 else regular_font
    c_o.number_format = num_fmt_qty if isinstance(item['nhap_thua'], float) and not item['nhap_thua'].is_integer() else num_fmt_int
    c_o.border = thin_border
    
    # Col P: Tổng thực xuất Net = (L{r}+M{r})-O{r}
    c_p = ws_master.cell(r, 16)
    c_p.value = f'=(L{r}+M{r})-O{r}'
    c_p.alignment = align_right
    c_p.font = bold_font
    c_p.number_format = num_fmt_qty
    c_p.border = thin_border
    
    # Col Q: Chênh lệch = P{r}-H{r}
    c_q = ws_master.cell(r, 17)
    c_q.value = f'=P{r}-H{r}'
    c_q.alignment = align_right
    c_q.font = bold_font
    c_q.number_format = num_fmt_qty
    c_q.border = thin_border
    
    # Col R: % Response = IF(H{r}>0, P{r}/H{r}, IF(P{r}>0, "Vượt", 1))
    c_r = ws_master.cell(r, 18)
    c_r.value = f'=IF(H{r}>0, P{r}/H{r}, IF(P{r}>0, "Vượt", 1))'
    c_r.alignment = align_right
    c_r.font = regular_font
    c_r.number_format = num_fmt_pct
    c_r.border = thin_border
    
    # Col S: Status
    c_stat = ws_master.cell(r, 19, item['status'])
    c_stat.alignment = align_center
    c_stat.font = bold_font
    c_stat.border = thin_border
    
    # Fill colors
    c_stat.fill = item['cell_fill']
    c_p.fill = item['cell_fill']
    c_q.fill = item['cell_fill']
        
    # Col T: Note
    c_note = ws_master.cell(r, 20, item['note'])
    c_note.alignment = align_left
    c_note.font = regular_font
    c_note.border = thin_border
    
    curr_master_row += 1

# Total row
total_row_idx = curr_master_row
ws_master.cell(total_row_idx, 1, 'Tổng cộng').font = bold_font
ws_master.cell(total_row_idx, 1).alignment = align_center

for c_idx in range(1, len(master_headers) + 1):
    c = ws_master.cell(total_row_idx, c_idx)
    c.fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    c.border = double_bottom_border
    c.font = bold_font

ws_master.cell(total_row_idx, 8, f'=SUM(H5:H{total_row_idx-1})').number_format = num_fmt_qty
ws_master.cell(total_row_idx, 12, f'=SUM(L5:L{total_row_idx-1})').number_format = num_fmt_qty
ws_master.cell(total_row_idx, 13, f'=SUM(M5:M{total_row_idx-1})').number_format = num_fmt_qty
ws_master.cell(total_row_idx, 14, f'=SUM(N5:N{total_row_idx-1})').number_format = num_fmt_qty
ws_master.cell(total_row_idx, 15, f'=SUM(O5:O{total_row_idx-1})').number_format = num_fmt_qty
ws_master.cell(total_row_idx, 16, f'=SUM(P5:P{total_row_idx-1})').number_format = num_fmt_qty
ws_master.cell(total_row_idx, 17, f'=SUM(Q5:Q{total_row_idx-1})').number_format = num_fmt_qty

ws_master.auto_filter.ref = f'A4:T{total_row_idx-1}'

# Save to destination file (target V2 first since previous file may be open)
dest_file_v2 = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units)_DOI_CHIEU_V2.xlsx'
wb.save(dest_file_v2)
print(f'Successfully saved upgraded reconciled workbook to: {dest_file_v2}')

# Attempt to save to _DOI_CHIEU.xlsx if closed
dest_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units)_DOI_CHIEU.xlsx'
try:
    wb.save(dest_file)
    print(f'Successfully updated: {dest_file}')
except PermissionError:
    print(f'Note: {dest_file} is currently open in Excel.')

# Attempt to save to original file if closed
orig_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units).xlsx'
try:
    wb.save(orig_file)
    print(f'Successfully overwritten original file: {orig_file}')
except PermissionError:
    print(f'Note: Original file is currently open in Excel.')
