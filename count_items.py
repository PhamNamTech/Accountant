import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

bom_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units).xlsx'
wb = openpyxl.load_workbook(bom_file, data_only=True)

thcp_map = {
    'HRM30A6001': 'DX1TNJ0023',
    'HRS20A6003': 'DX1TNJ0024',
    'DHRM38A6001': 'DX1TNJ0031',
    'SCR18A7501_MC_FSM': 'DX1TNJ0033',
    'GRU30A05-OF-L': 'DX1TNJ0039',
    'GRU30A05-OF-R': 'DX1TNJ0040',
}

for sname, thcp in thcp_map.items():
    ws = wb[sname]
    start_r = None
    for r in range(1, 200):
        if '1.2' in str(ws.cell(r, 1).value or ''):
            start_r = r
            break
            
    header_r = start_r + 1
    items = []
    for r in range(header_r + 1, header_r + 150):
        name = ws.cell(r, 3).value
        code = ws.cell(r, 4).value
        dvt = ws.cell(r, 6).value
        qty_proj = ws.cell(r, 11).value
        xk = ws.cell(r, 12).value
        note = ws.cell(r, 13).value
        
        # Check if name is filled
        if name and str(name).strip():
            items.append({
                'row': r,
                'stt': ws.cell(r, 2).value,
                'name': str(name).strip(),
                'code': str(code).strip() if code else '',
                'dvt': str(dvt).strip() if dvt else '',
                'qty_proj': qty_proj,
                'xk': xk,
                'note': str(note).strip() if note else ''
            })
            
    print(f'Sheet {sname} ({thcp}): {len(items)} valid material items in section 1.2')
