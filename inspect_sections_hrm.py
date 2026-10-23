import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

bom_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units).xlsx'
wb = openpyxl.load_workbook(bom_file, data_only=True)
ws = wb['HRM30A6001']

for r in range(1, 150):
    val = ws.cell(r, 1).value
    val2 = ws.cell(r, 2).value
    if val or val2:
        val_str = str(val or '') + ' ' + str(val2 or '')
        if any(k in val_str.upper() for k in ['PHẦN', 'MỤC', '1.', '2.', '3.', '4.', 'BẢNG', 'TỔNG']):
            print(f'Row {r:3d}: {val_str[:70]}')
