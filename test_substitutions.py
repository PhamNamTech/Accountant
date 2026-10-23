import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

bom_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units).xlsx'
wb = openpyxl.load_workbook(bom_file, data_only=True)

thcp_map = {
    'HRM30A6001': 'DX1TNJ0023',
    'HRS20A6003': 'DX1TNJ0024',
    'DHRM38A6001': 'DX1TNJ0031',
    'SCR18A7501_MC_FSM': 'DX1TNJ0033',
    'GRU30A05-OF-L': 'DX1TNJ0039',
    'GRU30A05-OF-R': 'DX1TNJ0040',
}

ws_main = wb['Đơn vị chính']
ws_sub = wb['đơn vị phụ']

misa_exports = {}
for r in range(6, ws_main.max_row):
    thcp = ws_main.cell(r, 21).value
    code = str(ws_main.cell(r, 3).value or '').strip()
    name = str(ws_main.cell(r, 4).value or '').strip()
    dvt = str(ws_main.cell(r, 10).value or '').strip()
    xuat_sl = ws_main.cell(r, 15).value or 0
    if thcp and code and xuat_sl > 0:
        if thcp not in misa_exports:
            misa_exports[thcp] = {}
        if code not in misa_exports[thcp]:
            misa_exports[thcp][code] = {'name': name, 'main_dvt': dvt, 'main_qty': 0, 'sub_dvt': None, 'sub_qty': 0}
        misa_exports[thcp][code]['main_qty'] += xuat_sl

for r in range(6, ws_sub.max_row):
    thcp = ws_sub.cell(r, 21).value
    code = str(ws_sub.cell(r, 3).value or '').strip()
    dvt = str(ws_sub.cell(r, 10).value or '').strip()
    xuat_sl = ws_sub.cell(r, 15).value or 0
    if thcp and code and xuat_sl > 0:
        if thcp in misa_exports and code in misa_exports[thcp]:
            misa_exports[thcp][code]['sub_dvt'] = dvt
            misa_exports[thcp][code]['sub_qty'] += xuat_sl

# Global substitution map
# (target_bom_code) -> [list of misa_codes that can substitute/merge into this bom code]
# We also do sheet-specific substitutions if needed
substitutions = {
    'CX1TBM0392': ['CX1HTP0610'], # Mặt bích tròn DN80
    'CX1TBM0393': ['CX1HTP0611'], # Mặt bích vuông DN80
    'CX1TBM0142': ['CX1HTP0548'], # Mặt bích tròn DN50
    'CX1TBM0365': ['CX1HTP0609'], # Mặt bích vuông DN50
    'CX1TBM0549': ['CX1HTP0549'], # Mặt bích tròn DN25
    'CX1KAN0158': ['CX1MAX0691'], # Xích 35-1R (merge)
    'CX1KAN0157': ['CX1NOA0470', 'CX1STA0318', 'CX1DID0168'], # Xích 40-1R (merge)
    'CX1NOA0470': ['CX1STA0318'], # Xích 40-1R
    'CX1NOA0300': ['CX1NOA0681'], # Long đền M14x30 (1mm + 2mm)
    'CX1NOA0277': ['CX1NOA0682'], # Long đền M16x30 (1mm + 2mm)
    'CX1NOA0455': ['CX1NOA0680'], # Long đền M8x19 (1mm + 2mm)
    'CX1NOA0149': ['CX1NOA0537'], # Nhựa POM DN40 (merge)
    'CX1NOA0378': ['CX1NOA0377'], # Ron Teflon DN80
    'CX1NTN0250': ['CX1PBC0429'], # Bạc đạn 6808
    'CX1SKF0111': ['CX1UKB0692'], # Bạc đạn 6206
    'CX1NOA0130': ['CX1MAX0690'], # Nhông xích 40B-12T
    'CX1KLO0254': ['CX1KLO0600'], # Nắp hộp giảm tốc
    'CX1KLO0422': ['CX1KLO0596'], # MC Khâu 2
    'CX1VVU0584': ['CX1VVU0585'], # Pat điều khiển
    'CX1NOA0468': ['CX1NOA0704'], # Bulong M8x15 / M8x16
    'CX1NOA0640': ['CX1NOA0450'], # Bulong M5x15 / M5x16
    '450': ['CX1NOA0450'],
    '700': ['CX1NOA0700'],
    '702': ['CX1NOA0702'],
    '703': ['CX1NOA0703'],
    '704': ['CX1NOA0704'],
    '309': ['CX1NOA0309'],
    'CX1KLO0605': ['CX1KLO0605'],
    'CX1RBI0445': ['CX1FBJ0705'], # Bạc đạn 696 ZZ
    'CX1DPI0445': ['CX1FBJ0705', 'CX1RBI0564'],
    'CX1KLO0413': ['CX1HTP0628', 'CX1HTP0629'], # Lưỡi gài / chốt giữ
    'CX1NOA0105': ['CX1NOA0013'], # Cây ty tròn đặc 14mm / 18mm
    'CX1KLO0420': ['CX1SFM0683'], # Khớp xoay SVE12S02
    'CX1PLY0447': ['CX1PLY0698'], # Lò xo cài cóc
    'CX1DWH0503': ['CX1DWH0502'], # Lò xo xoắn cuộn
    'CX1LKB0693': ['CX1LKB0694'], # Gối đỡ UCP212
}

print('Testing matching with substitution map...')

for sname, thcp in thcp_map.items():
    ws = wb[sname]
    start_r = None
    for r in range(1, 200):
        if '1.2' in str(ws.cell(r, 1).value or ''):
            start_r = r
            break
            
    header_r = start_r + 1
    misa_map = misa_exports.get(thcp, {})
    misa_claimed = set()
    
    print(f'=== {sname} ({thcp}) ===')
    matched_count = 0
    total_bom_items = 0
    
    for r in range(header_r + 1, ws.max_row + 1):
        v_a = ws.cell(r, 1).value
        v_b = ws.cell(r, 2).value
        v_c = ws.cell(r, 3).value
        v_d = ws.cell(r, 4).value
        v_f = ws.cell(r, 6).value
        v_k = ws.cell(r, 11).value
        v_l = ws.cell(r, 12).value
        v_m = ws.cell(r, 13).value
        
        if v_a and any(k in str(v_a).upper() for k in ['PHẦN 2', '2.', 'TỔNG CỘNG']):
            break
        if not (v_c or v_d):
            continue
            
        total_bom_items += 1
        c_code = str(v_d).strip() if v_d else ''
        c_dvt = str(v_f).strip() if v_f else ''
        
        # Check direct and substitutes
        codes_to_check = [c_code]
        if c_code in substitutions:
            codes_to_check.extend(substitutions[c_code])
        # Short digit codes e.g. 700 -> CX1NOA0700
        if c_code.isdigit() and len(c_code) == 3:
            codes_to_check.append(f'CX1NOA0{c_code}')
            
        # Also check name matching if code empty or #N/A
        if not c_code or c_code in ('#N/A', 'None'):
            # try name
            for mc, mv in misa_map.items():
                if mc not in misa_claimed:
                    if mv['name'].lower() in str(v_c).lower() or str(v_c).lower() in mv['name'].lower():
                        codes_to_check.append(mc)
                        
        found_misa = []
        for ck in codes_to_check:
            if ck in misa_map and ck not in misa_claimed:
                found_misa.append((ck, misa_map[ck]))
                misa_claimed.add(ck)
                
        if found_misa:
            matched_count += 1
            # print matched info
            m_codes_str = '+'.join([c for c, _ in found_misa])
            # calculate qty in BOM unit
            # if BOM unit is mm and MISA is m -> * 1000
            total_qty_bom = 0
            for c_found, m_obj in found_misa:
                # determine if main or sub matches BOM dvt
                if c_dvt.lower() in ['mm', 'm']:
                    if m_obj['main_dvt'].lower() in ['m', 'mét']:
                        total_qty_bom += m_obj['main_qty'] * 1000 if c_dvt.lower() == 'mm' else m_obj['main_qty']
                    elif m_obj['sub_dvt'] and m_obj['sub_dvt'].lower() in ['m', 'mét']:
                        total_qty_bom += m_obj['sub_qty'] * 1000 if c_dvt.lower() == 'mm' else m_obj['sub_qty']
                elif c_dvt.lower() in ['que', 'cây'] and m_obj['main_dvt'].lower() == 'kg' and m_obj['sub_dvt'] and m_obj['sub_dvt'].lower() in ['que', 'cây']:
                    total_qty_bom += m_obj['sub_qty']
                elif c_dvt.lower() in ['cái', 'con', 'bộ', 'chai', 'lon', 'hộp', 'đôi', 'gói', 'tấm', 'cặp', 'mũi']:
                    if m_obj['main_dvt'].lower() in ['cái', 'cái', 'con', 'bộ', 'chai', 'lon', 'hộp', 'đôi', 'gói', 'tấm', 'cặp', 'mũi']:
                        total_qty_bom += m_obj['main_qty']
                    elif m_obj['sub_dvt'] and m_obj['sub_dvt'].lower() in ['cái', 'cái', 'con', 'bộ', 'chai', 'lon', 'hộp', 'đôi', 'gói', 'tấm', 'cặp', 'mũi']:
                        total_qty_bom += m_obj['sub_qty']
                    else:
                        total_qty_bom += m_obj['main_qty']
                else:
                    total_qty_bom += m_obj['main_qty']
                    
            print(f'  Row {r:3d} | BOM: {c_code:12s} ({c_dvt:5s}) | ĐM: {str(v_k):>8} | File XK: {str(v_l):>8} | Calc MISA: {total_qty_bom:>8} ({m_codes_str})')
            
    unclaimed = set(misa_map.keys()) - misa_claimed
    print(f'-> Matched: {matched_count}/{total_bom_items} BOM rows. Unclaimed MISA codes ({len(unclaimed)}): {unclaimed}\n')
