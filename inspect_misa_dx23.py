import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

bom_file = r'F:\Accountant\data\BOM_FSM Master (Goi 2 10units).xlsx'
wb = openpyxl.load_workbook(bom_file, data_only=True)

ws_main = wb['Đơn vị chính']
ws_sub = wb['đơn vị phụ']

print('=== MISA Xuất kho for DX1TNJ0023 (HRM30A6001) in Đơn vị chính ===')
misa_main_23 = []
for r in range(6, ws_main.max_row):
    thcp = ws_main.cell(r, 21).value
    if thcp == 'DX1TNJ0023':
        code = ws_main.cell(r, 3).value
        name = ws_main.cell(r, 4).value
        dvt = ws_main.cell(r, 10).value
        xuat_sl = ws_main.cell(r, 15).value
        misa_main_23.append((r, code, name, dvt, xuat_sl))
        print(f'Row {r:4d} | Code: {code:12s} | ĐVT: {dvt:6s} | Xuất SL: {xuat_sl:>8} | Tên: {name[:35]}')

print(f'\nTotal export transactions in Đơn vị chính for DX1TNJ0023: {len(misa_main_23)}')

print('\n=== MISA Xuất kho for DX1TNJ0023 (HRM30A6001) in đơn vị phụ ===')
misa_sub_23 = []
for r in range(6, ws_sub.max_row):
    thcp = ws_sub.cell(r, 21).value
    if thcp == 'DX1TNJ0023':
        code = ws_sub.cell(r, 3).value
        name = ws_sub.cell(r, 4).value
        dvt = ws_sub.cell(r, 10).value
        xuat_sl = ws_sub.cell(r, 15).value
        misa_sub_23.append((r, code, name, dvt, xuat_sl))
        print(f'Row {r:4d} | Code: {code:12s} | ĐVT: {dvt:6s} | Xuất SL: {xuat_sl:>8} | Tên: {name[:35]}')

print(f'\nTotal export transactions in đơn vị phụ for DX1TNJ0023: {len(misa_sub_23)}')
